# app/views.py
from django.contrib.auth.decorators import login_required
import json
import os
import zipfile
import pandas as pd
from io import BytesIO
from django.shortcuts import render, redirect
from django.http import HttpResponse, HttpResponseBadRequest, JsonResponse
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from django.db.models import Count
from .forms import DosenForm 
from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.forms import PasswordChangeForm
from django.contrib.auth import update_session_auth_hash
from .forms import GantiAkunForm  # Pastikan form ini sudah ada di forms.py
from django.contrib.auth.models import User
from django.http import HttpResponse

# IMPORT MODEL RELASIONAL (Pastikan models.py sudah ada kolom sejak_2020)
from .models import Dosen, Publikasi, StatistikSitasi
from . import scraping

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOSEN_JSON = os.path.join(BASE_DIR, "app", "dosen.json")

def load_dosen_list():
    try:
        with open(DOSEN_JSON, "r", encoding="utf-8") as f:
            return json.load(f)
    except:
        return []

# =========================================================
#  BAGIAN 1: HELPER FUNCTIONS (LOGIKA DATABASE & EXCEL)
# =========================================================

def save_dosen_to_db(author_id, data_pack):
    """
    Fungsi Penyimpanan dengan LOGIKA HYBRID (Indo/Inggris)
    Mengadopsi logika scraping lama kamu agar data tidak 0.
    """
    author_json = data_pack.get("author_json", {})
    
    # 1. AMBIL INFO UTAMA
    info = author_json.get("author", {}) 
    
    # 2. AMBIL ARTIKEL
    all_articles = data_pack.get("all_articles", [])
    if not all_articles:
        all_articles = data_pack.get("articles", [])
    
    # -----------------------------------------------------------
    # [LOGIKA BARU - HYBRID] Adaptasi dari kode lama kamu
    # -----------------------------------------------------------
    cited_by = author_json.get("cited_by", {})
    table = cited_by.get("table", [])
    
    # Inisialisasi 0 semua
    cit_all = 0; cit_since = 0
    h_all = 0; h_since = 0
    i10_all = 0; i10_since = 0

    # Fungsi kecil untuk bongkar baris tabel (Bisa baca 'kutipan' atau 'citations')
    def extract_stats_row(row_data):
        if not row_data: return 0, 0
        # row_data isinya bisa: {'citations': {'all': 121, 'since_2020': 121}} 
        # ATAU {'kutipan': {'all': 121, 'sejak_2020': 121}}
        
        # Ambil isi dalam dictionary (abaikan kunci luarnya apa)
        inner_values = list(row_data.values())[0] if row_data else {}
        
        val_all = inner_values.get("all", 0)
        val_since = 0
        
        # Cari kunci yang depannya 'since_' atau 'sejak_'
        for k, v in inner_values.items():
            if k.startswith("since_") or k.startswith("sejak_"):
                val_since = v
                break # Ketemu, langsung stop
        return val_all, val_since

    # Terapkan ke Baris 1, 2, dan 3
    if len(table) > 0:
        cit_all, cit_since = extract_stats_row(table[0]) # Baris 1: Sitasi
    if len(table) > 1:
        h_all, h_since = extract_stats_row(table[1])     # Baris 2: H-index
    if len(table) > 2:
        i10_all, i10_since = extract_stats_row(table[2]) # Baris 3: i10-index

    # 3. OLAH BIDANG MINAT
    interests_raw = info.get("interests", [])
    if isinstance(interests_raw, list):
        bidang_str = ", ".join([i.get("title", "") for i in interests_raw if isinstance(i, dict)])
    else:
        bidang_str = "-"

    # 4. SIMPAN KE DATABASE (RELASIONAL)
    dosen_obj, created = Dosen.objects.update_or_create(
        author_id=author_id,
        defaults={
            'nama': info.get("name"),
            'afiliasi': info.get("affiliations"),
            'foto_url': info.get("thumbnail"), 
            'email': info.get("email"),       
            'bidang_minat': bidang_str,       
            
            # ISI STATISTIK LENGKAP
            'total_sitasi': int(cit_all),
            'h_index': int(h_all),
            'i10_index': int(i10_all),
            
            'total_sitasi_sejak_2020': int(cit_since),
            'h_index_sejak_2020': int(h_since),
            'i10_index_sejak_2020': int(i10_since),
        }
    )

    # 5. SIMPAN PUBLIKASI
    Publikasi.objects.filter(dosen=dosen_obj).delete()
    list_pub_obj = []
    for art in all_articles:
        cit_raw = art.get("cited_by", {}).get("value", 0)
        try: cit_val = int(cit_raw)
        except: cit_val = 0
    
        pub = Publikasi(
            dosen=dosen_obj,
            judul=art.get("title", "No Title")[:499], 
            link_jurnal=art.get("link") or art.get("external_link"),
            authors=art.get("authors", "-"),
            nama_jurnal=art.get("publication", "-"),
            tahun=str(art.get("year", "-")),
            jumlah_sitasi=cit_val
        )
        list_pub_obj.append(pub)
    Publikasi.objects.bulk_create(list_pub_obj)

    # 6. SIMPAN GRAFIK
    StatistikSitasi.objects.filter(dosen=dosen_obj).delete()
    graph_data = cited_by.get("graph", [])
    list_stat_obj = []
    for g in graph_data:
        try:
            stat = StatistikSitasi(
                dosen=dosen_obj,
                tahun=int(g.get("year")),
                jumlah=int(g.get("citations"))
            )
            list_stat_obj.append(stat)
        except: continue
    StatistikSitasi.objects.bulk_create(list_stat_obj)
    
    return dosen_obj

def generate_excel_bytes_from_model(dosen_obj):
    """
    Generate Excel dari Model Relasional (Dosen -> Publikasi)
    """
    # 1. Sheet Profil
    df_info = pd.DataFrame([{
        "Nama Dosen": dosen_obj.nama,
        "Afiliasi": dosen_obj.afiliasi,
        "Email": dosen_obj.email, 
        "Bidang Minat": dosen_obj.bidang_minat,
        "Total Kutipan (Semua)": dosen_obj.total_sitasi,
        "Total Kutipan (Sejak 2020)": dosen_obj.total_sitasi_sejak_2020,
        "h-index (Semua)": dosen_obj.h_index,
        "h-index (Sejak 2020)": dosen_obj.h_index_sejak_2020,
        "i10-index (Semua)": dosen_obj.i10_index,
        "i10-index (Sejak 2020)": dosen_obj.i10_index_sejak_2020,
        "Terakhir Update": dosen_obj.last_updated.strftime("%Y-%m-%d %H:%M")
    }])

    # 2. Sheet Publikasi
    pub_data = []
    for p in dosen_obj.publikasi.all():
        pub_data.append({
            "Judul": p.judul,
            "Penulis": p.authors,
            "Tahun": p.tahun,
            "Jurnal": p.nama_jurnal,
            "Sitasi": p.jumlah_sitasi,
            "Link": p.link_jurnal
        })
    df_pub = pd.DataFrame(pub_data)

    # 3. Sheet Grafik
    graph_data = []
    for s in dosen_obj.statistik.all().order_by('tahun'):
        graph_data.append({"Tahun": s.tahun, "Jumlah": s.jumlah})
    df_graph = pd.DataFrame(graph_data)

    # 4. Tulis ke Excel
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_info.to_excel(writer, sheet_name="Profil", index=False)
        df_pub.to_excel(writer, sheet_name="Publikasi", index=False)
        df_graph.to_excel(writer, sheet_name="Grafik", index=False)
    
    # Tambahkan Chart
    buffer.seek(0)
    wb = load_workbook(buffer)
    if not df_graph.empty:
        ws = wb["Grafik"]
        chart = BarChart()
        chart.title = "Statistik Sitasi"
        data_ref = Reference(ws, min_col=2, min_row=1, max_row=len(df_graph)+1)
        cats_ref = Reference(ws, min_col=1, min_row=2, max_row=len(df_graph)+1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        ws.add_chart(chart, "D5")
        
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out

# =========================================================
#  BAGIAN 2: VIEWS UTAMA
# =========================================================

@login_required(login_url='login')
def dashboard_view(request):
    # --- LOGIKA BARU: Hitung Real-time dari Database ---
    jml_feb = Dosen.objects.filter(fakultas='FEB').count()
    jml_ftd = Dosen.objects.filter(fakultas='FTD').count()
    
    stats = {
        'fakultas': 2, 
        'prodi': 5,
        'dosen_ftd': jml_ftd, # Ini sekarang otomatis berubah
        'dosen_feb': jml_feb  # Ini sekarang otomatis berubah
    }
    return render(request, "app/dashboard.html", stats)

@login_required(login_url='login')
def author_view(request):
    # GANTI load_dosen_list() DENGAN INI:
    dosen_db = Dosen.objects.all().order_by('nama')
    dosen_list = [{'id': d.author_id, 'nama': d.nama} for d in dosen_db]
    selected_id = request.GET.get("id")
    export = request.GET.get("export")
    do_sync = request.GET.get("sync") == "true"

    if do_sync and selected_id:
        try:
            print(f"🔄 SINGLE SYNC: {selected_id}...")
            data_pack = scraping.get_full_data_via_serpapi(selected_id)
            save_dosen_to_db(selected_id, data_pack)
            return redirect(f"/data-dosen/?id={selected_id}")
        except Exception as e: print(f"❌ Error Sync: {e}")

    if export == "excel" and selected_id:
        try:
            dosen_obj = Dosen.objects.get(author_id=selected_id)
            excel_bytes = generate_excel_bytes_from_model(dosen_obj)
            resp = HttpResponse(excel_bytes.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            resp['Content-Disposition'] = f'attachment; filename="{dosen_obj.nama}.xlsx"'
            return resp
        except Dosen.DoesNotExist:
            return HttpResponse("Data belum ada, silakan update dulu.", status=404)

    context = {"dosen_list": dosen_list, "selected_id": selected_id}
    if selected_id:
        try:
            dosen = Dosen.objects.get(author_id=selected_id)
            publications = dosen.publikasi.all().order_by('-tahun')[:20]
            
            labels = []
            values = []
            for s in dosen.statistik.all().order_by('tahun'):
                labels.append(s.tahun)
                values.append(s.jumlah)
            
            context.update({
                "profile": dosen, 
                "publications": publications,
                "labels_json": json.dumps(labels),
                "values_json": json.dumps(values)
            })
        except Dosen.DoesNotExist:
            pass

    return render(request, "app/index.html", context)

# @login_required(login_url='login')
# def update_all_data(request):
#     dosen_list = load_dosen_list()
#     success_count = 0
#     print("\n🚀 MEMULAI UPDATE MASSAL RELASIONAL...")
#     for i, dosen in enumerate(dosen_list):
#         d_id = dosen['id']
#         print(f"[{i+1}/{len(dosen_list)}] Updating DB: {dosen['nama']}...")
#         try:
#             data_pack = scraping.get_full_data_via_serpapi(d_id)
#             save_dosen_to_db(d_id, data_pack)
#             success_count += 1
#         except Exception as e:
#             print(f"❌ Gagal update {dosen['nama']}: {e}")
#     return JsonResponse({"status": "success", "message": f"Berhasil memperbarui {success_count} data dosen!"})


# --- TAMBAHKAN VARIABEL GLOBAL INI ---
GLOBAL_STOP_UPDATE = False

@login_required(login_url='login')
def cancel_update_view(request):
    """Fungsi kecil untuk menarik rem tangan (membatalkan loop)"""
    global GLOBAL_STOP_UPDATE
    GLOBAL_STOP_UPDATE = True
    return JsonResponse({"message": "Stop signal received"})

@login_required(login_url='login')
def update_all_data(request):
    global GLOBAL_STOP_UPDATE
    GLOBAL_STOP_UPDATE = False # Lepas rem saat proses baru dimulai
    
    dosen_list = load_dosen_list()
    success_count = 0
    print("\n🚀 MEMULAI UPDATE MASSAL RELASIONAL...")
    
    for i, dosen in enumerate(dosen_list):
        # --- CEK REM TANGAN DI SINI ---
        if GLOBAL_STOP_UPDATE:
            print("🛑 UPDATE MASSAL DIBATALKAN OLEH USER DI TENGAH JALAN!")
            break # Hentikan perulangan secara paksa!
            
        d_id = dosen['id']
        print(f"[{i+1}/{len(dosen_list)}] Updating DB: {dosen['nama']}...")
        try:
            data_pack = scraping.get_full_data_via_serpapi(d_id)
            save_dosen_to_db(d_id, data_pack)
            success_count += 1
        except Exception as e:
            print(f"❌ Gagal update {dosen['nama']}: {e}")
            
    return JsonResponse({"status": "success", "message": f"Berhasil memperbarui {success_count} data dosen!"})

@login_required(login_url='login')
def download_all_excel(request):
    dosen_list = load_dosen_list()
    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for dosen in dosen_list:
            d_id = dosen['id']
            try:
                dosen_obj = Dosen.objects.get(author_id=d_id)
                excel_bytes = generate_excel_bytes_from_model(dosen_obj)
                zip_file.writestr(f"{dosen_obj.nama}.xlsx", excel_bytes.getvalue())
            except Dosen.DoesNotExist: pass
    zip_buffer.seek(0)
    response = HttpResponse(zip_buffer.read(), content_type="application/zip")
    response['Content-Disposition'] = 'attachment; filename="Semua_Data_Dosen.zip"'
    return response

@login_required(login_url='login')
def dosen_list_view(request):
    dosen_list = load_dosen_list()
    cached_ids = Dosen.objects.values_list('author_id', flat=True)
    for d in dosen_list:
        if d['id'] in cached_ids:
            d['status'] = "Tersimpan"; d['class'] = "success"
        else:
            d['status'] = "Belum Ada"; d['class'] = "secondary"
    context = {'dosen_list': dosen_list, 'total_dosen': len(dosen_list), 'total_cached': len(cached_ids)}
    return render(request, "app/dosen_list.html", context)

@login_required(login_url='login')
@login_required(login_url='login')
def leaderboard_view(request):
    # Kita pakai .annotate() untuk menghitung jumlah publikasi (count) secara otomatis
    ranking_list = Dosen.objects.annotate(jumlah_publikasi=Count('publikasi')).order_by('-total_sitasi')

    context = {
        'ranking_list': ranking_list,
        'total_data': ranking_list.count()
    }
    return render(request, "app/leaderboard.html", context)

@login_required(login_url='login')
def detail_dosen_view(request, author_id):
    try:
        dosen = Dosen.objects.get(author_id=author_id)
        publications = dosen.publikasi.all().order_by('-tahun')
        context = {
            'profile': dosen,
            'publications': publications,
            'total_cit': dosen.total_sitasi,
            'total_pub': publications.count(),
        }
        return render(request, "app/detail_dosen.html", context)
    except Dosen.DoesNotExist:
        return HttpResponse("Data dosen ini belum ada di database.", status=404)
    
# =========================================================
#  FITUR BARU: KELOLA DOSEN (TAMBAH & HAPUS)
# =========================================================

@login_required(login_url='login')
# def kelola_dosen_view(request):
#     # 1. Logika Tambah Dosen (Saat tombol Simpan ditekan)
#     if request.method == 'POST':
#         form = DosenForm(request.POST)
#         if form.is_valid():
#             form.save()
#             messages.success(request, "Dosen berhasil ditambahkan! Silakan update datanya.")
#             return redirect('kelola_dosen')
#     else:
#         form = DosenForm()

#     # 2. Logika Menampilkan List Dosen di Tabel
#     list_dosen = Dosen.objects.all().order_by('-last_updated')

#     context = {
#         'form': form,
#         'list_dosen': list_dosen
#     }
#     return render(request, "app/kelola_dosen.html", context)

@login_required(login_url='login')
def kelola_dosen_view(request):
    # 1. Logika Tambah Dosen
    if request.method == 'POST':
        form = DosenForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Dosen berhasil ditambahkan! Silakan update datanya.")
            return redirect('kelola_dosen')
    else:
        form = DosenForm()

    # 2. Logika Menampilkan List Dosen di Tabel
    list_dosen = Dosen.objects.all().order_by('-last_updated')

    # --- KODE BARU UNTUK CEK KUOTA API ---
    import requests
    from django.conf import settings
    
    sisa_kuota = "-"
    try:
        url = f"https://serpapi.com/account.json?api_key={settings.SERPAPI_KEY}"
        response = requests.get(url, timeout=5)
        sisa_kuota = response.json().get("plan_searches_left", "Error")
    except:
        sisa_kuota = "Gagal"
    # -------------------------------------

    context = {
        'form': form,
        'list_dosen': list_dosen,
        'sisa_kuota': sisa_kuota 
    }
    return render(request, "app/kelola_dosen.html", context)
#########################################################################

@login_required(login_url='login')
def hapus_dosen_view(request, author_id):
    # Logika Hapus Dosen
    try:
        dosen = Dosen.objects.get(author_id=author_id)
        dosen.delete() # Hapus dari database (otomatis hilang di dashboard & peringkat)
        messages.success(request, "Data dosen berhasil dihapus.")
    except Dosen.DoesNotExist:
        messages.error(request, "Dosen tidak ditemukan.")
    
    return redirect('kelola_dosen')



@login_required(login_url='login')
def edit_dosen_view(request, author_id):
    # Ambil data lama
    dosen_lama = get_object_or_404(Dosen, author_id=author_id)
    
    if request.method == 'POST':
        form = DosenForm(request.POST, instance=dosen_lama)
        if form.is_valid():
            # Cek apakah ID berubah?
            new_id = form.cleaned_data['author_id']
            old_id = dosen_lama.author_id
            
            if new_id != old_id:
                # 1. Simpan Dosen Baru (ID Baru)
                dosen_baru = form.save(commit=False)
                dosen_baru.author_id = new_id
                dosen_baru.save() # Data baru tercipta
                
                # 2. PINDAHKAN ANAK-ANAKNYA (Publikasi & Statistik) ke Bapak Baru
                Publikasi.objects.filter(dosen_id=old_id).update(dosen_id=new_id)
                StatistikSitasi.objects.filter(dosen_id=old_id).update(dosen_id=new_id)
                
                # 3. Hapus Bapak Lama
                Dosen.objects.get(author_id=old_id).delete()
                
                messages.success(request, f"Data berhasil diedit! ID berubah dari {old_id} ke {new_id}.")
            else:
                # Kalau ID tidak berubah, simpan biasa
                form.save()
                messages.success(request, "Data berhasil diupdate!")

            return redirect('kelola_dosen')
    else:
        form = DosenForm(instance=dosen_lama)

    return render(request, "app/edit_dosen.html", {'form': form, 'dosen': dosen_lama})

@login_required(login_url='login')
def ganti_password_view(request):
    if request.method == 'POST':
        form = PasswordChangeForm(request.user, request.POST)
        if form.is_valid():
            user = form.save()
            # Mencegah user ter-logout otomatis setelah ganti password
            update_session_auth_hash(request, user)
            messages.success(request, "Password berhasil diperbarui!")
            return redirect('dashboard')
        else:
            messages.error(request, "Silakan perbaiki kesalahan di bawah.")
    else:
        form = PasswordChangeForm(request.user)
    
    return render(request, 'app/ganti_password.html', {'form': form})

@login_required(login_url='login')
def ganti_akun_view(request):
    if request.method == 'POST':
        form = GantiAkunForm(request.POST, instance=request.user)
        if form.is_valid():
            user = form.save(commit=False)
            new_username = form.cleaned_data.get('username')
            
            # Cek jika username sudah dipakai akun lain
            from django.contrib.auth.models import User
            if User.objects.filter(username=new_username).exclude(pk=request.user.pk).exists():
                messages.error(request, "Username tersebut sudah digunakan akun lain!")
                return redirect('ganti_akun')

            user.username = new_username
            
            # Jika password baru diisi, perbarui passwordnya
            p1 = form.cleaned_data.get('password_baru1')
            if p1:
                user.set_password(p1)
                update_session_auth_hash(request, user) # Mencegah user ter-logout otomatis
                
            user.save()
            messages.success(request, "Akun (Username/Password) berhasil diperbarui!")
            return redirect('dashboard')
    else:
        form = GantiAkunForm(instance=request.user)

    return render(request, 'app/ganti_akun.html', {'form': form})

def buat_admin_sementara(request):
    if not User.objects.filter(username='adminbaru').exists():
        User.objects.create_superuser('adminbaru', 'admin@web.com', 'Rahasia123!')
        return HttpResponse("BERHASIL! Username: adminbaru | Password: Rahasia123!")
    return HttpResponse("Akun sudah ada!")