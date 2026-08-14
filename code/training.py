# import requests
# import pandas as pd
# import time
# import openpyxl
# from openpyxl.chart import BarChart, Reference

# # ============================
# # 🔑 API KEY SERPAPI
# # ============================
# SERPAPI_API_KEY = "1eb991a416303d0888fee987b55b42ca60ea7fd6d75f1ee0132782515a461d36"

# # ============================
# # 📘 Ambil data profil dosen
# # ============================
# def get_author_info(user_id):
#     params = {
#         "engine": "google_scholar_author",
#         "author_id": user_id,
#         "hl": "id",
#         "api_key": SERPAPI_API_KEY
#     }
#     res = requests.get("https://serpapi.com/search", params=params)
#     return res.json()

# # ============================
# # 📗 Ambil daftar publikasi dosen
# # ============================
# def get_publications_for_author(user_id):
#     publications = []
#     start = 0

#     while True:
#         params = {
#             "engine": "google_scholar_author",
#             "author_id": user_id,
#             "hl": "id",
#             "api_key": SERPAPI_API_KEY,
#             "start": start
#         }
#         res = requests.get("https://serpapi.com/search", params=params)
#         data = res.json()
#         pubs = data.get("articles", [])
#         if not pubs:
#             break

#         for pub in pubs:
#             link = pub.get("external_link") or pub.get("link")
#             publications.append({
#                 "Judul Penelitian": pub.get("title"),
#                 "Link Publikasi": link,
#                 "Tipe Link": "PDF" if link and link.endswith(".pdf") else "Halaman Web",
#                 "Pengarang": pub.get("authors"),
#                 "Jumlah Sitasi": pub.get("cited_by", {}).get("value", 0),
#                 "Tahun Terbit": pub.get("year"),
#                 "Nama Jurnal": pub.get("publication")
#             })

#         start += len(pubs)
#         time.sleep(1)

#     return publications

# # ============================
# # 📊 Ambil statistik kutipan (total, h-index, i10-index)
# # ============================
# def extract_citation_stats(author_info):
#     cited_by = author_info.get("cited_by", {})
#     table = cited_by.get("table", [])

#     total_kutipan = 0
#     h_index = 0
#     i10_index = 0

#     try:
#         total_kutipan = table[0].get("kutipan", {}).get("all", 0)
#         h_index = table[1].get("indeks_h", {}).get("all", 0)
#         i10_index = table[2].get("indeks_i10", {}).get("all", 0)
#     except Exception:
#         pass

#     return total_kutipan, h_index, i10_index

# # ============================
# # 📈 Ambil data grafik kutipan per tahun
# # ============================
# def extract_graph_data(author_info):
#     cited_by = author_info.get("cited_by", {})
#     graph_data = cited_by.get("graph", [])
#     return [{"Tahun": g.get("year"), "Jumlah Kutipan": g.get("citations")} for g in graph_data]

# # ============================
# # 👨‍🏫 Daftar Dosen
# # ============================
# dosen_list = [
#     {"id": "3ZEzr0EAAAAJ", "nama": "Handry Rochmad Dwi Happy"},
#     {"id": "IabXZ4wAAAAJ", "nama": "RISA SANTOSO"},
# ]

# # ============================
# # 🚀 Proses Setiap Dosen
# # ============================
# for dosen in dosen_list:
#     print(f"\n📄 Mengambil data untuk {dosen['nama']} (ID: {dosen['id']})...")

#     author_info = get_author_info(dosen["id"])
#     total_kutipan, h_index, i10_index = extract_citation_stats(author_info)
#     publications = get_publications_for_author(dosen["id"])
#     graph_data = extract_graph_data(author_info)

#     if not publications:
#         print(f"⚠️ Tidak ada publikasi ditemukan untuk {dosen['nama']}.")
#         continue

#     # DataFrame profil
#     df_info = pd.DataFrame({
#         "Nama Dosen": [author_info.get("author", {}).get("name", dosen["nama"])],
#         "Afiliasi": [author_info.get("author", {}).get("affiliations", "")],
#         "Email": [author_info.get("author", {}).get("email", "")],
#         "Bidang": [", ".join([i.get("title", "") for i in author_info.get("author", {}).get("interests", [])])],
#         "Total Kutipan": [total_kutipan],
#         "h-index": [h_index],
#         "i10-index": [i10_index]
#     })

#     df_pub = pd.DataFrame(publications)
#     df_graph = pd.DataFrame(graph_data)

#     # Simpan ke Excel
#     file_name = f"{dosen['nama']}_publikasi.xlsx"
#     with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
#         df_info.to_excel(writer, index=False, sheet_name="Profil Dosen")
#         df_pub.to_excel(writer, index=False, sheet_name="Daftar Publikasi")
#         if not df_graph.empty:
#             df_graph.to_excel(writer, index=False, sheet_name="Grafik Kutipan")

#     # ============================
#     # 📊 Tambahkan grafik batang di Excel
#     # ============================
#     if not df_graph.empty:
#         wb = openpyxl.load_workbook(file_name)
#         ws = wb["Grafik Kutipan"]

#         chart = BarChart()
#         chart.title = "Grafik Kutipan per Tahun"
#         chart.x_axis.title = "Tahun"
#         chart.y_axis.title = "Jumlah Kutipan"

#         data = Reference(ws, min_col=2, min_row=1, max_row=len(df_graph) + 1)
#         categories = Reference(ws, min_col=1, min_row=2, max_row=len(df_graph) + 1)
#         chart.add_data(data, titles_from_data=True)
#         chart.set_categories(categories)

#         ws.add_chart(chart, "E2")
#         wb.save(file_name)

#     print(f"✅ Data {dosen['nama']} berhasil disimpan ke file: {file_name}")
#     print(f"📊 Total Kutipan: {total_kutipan} | h-index: {h_index} | i10-index: {i10_index}")
################################################################################################################

# #YANG KE 2
# # save as export_author_with_citations.py
# import requests
# import pandas as pd
# import time
# import openpyxl

# SERPAPI_API_KEY = "1eb991a416303d0888fee987b55b42ca60ea7fd6d75f1ee0132782515a461d36"

# def get_author_info(user_id):
#     params = {
#         "engine": "google_scholar_author",
#         "author_id": user_id,
#         "hl": "id",
#         "api_key": SERPAPI_API_KEY
#     }
#     res = requests.get("https://serpapi.com/search", params=params)
#     res.raise_for_status()
#     return res.json()

# def get_publications_for_author(user_id):
#     publications = []
#     start = 0
#     while True:
#         params = {
#             "engine": "google_scholar_author",
#             "author_id": user_id,
#             "hl": "id",
#             "api_key": SERPAPI_API_KEY,
#             "start": start
#         }
#         res = requests.get("https://serpapi.com/search", params=params)
#         res.raise_for_status()
#         data = res.json()
#         pubs = data.get("articles", []) or data.get("articles", [])
#         if not pubs:
#             break

#         for pub in pubs:
#             link = pub.get("external_link") or pub.get("link")
#             publications.append({
#                 "Judul Penelitian": pub.get("title"),
#                 "Link Publikasi": link,
#                 "Tipe Link": "PDF" if link and isinstance(link, str) and link.lower().endswith(".pdf") else "Halaman Web",
#                 "Pengarang": pub.get("authors"),
#                 "Jumlah Sitasi": pub.get("cited_by", {}).get("value", 0) if pub.get("cited_by") else 0,
#                 "Tahun Terbit": pub.get("year"),
#                 "Nama Jurnal": pub.get("publication")
#             })
#         start += len(pubs)
#         time.sleep(1)  # jeda untuk hindari rate limit

#     return publications

# def extract_citation_stats(author_info):
#     """
#     Mengembalikan tuple:
#     (total_all, total_since2020, h_all, h_since2020, i10_all, i10_since2020)
#     Aman jika struktur JSON tidak lengkap.
#     """
#     cited_by = author_info.get("cited_by", {}) or {}
#     table = cited_by.get("table", []) if isinstance(cited_by.get("table", []), list) else []

#     # default
#     total_all = total_since = 0
#     h_all = h_since = 0
#     i10_all = i10_since = 0

#     # Table biasanya list: [ { "kutipan": {...} }, { "indeks_h": {...} }, { "indeks_i10": {...} } ]
#     try:
#         if len(table) >= 1:
#             total_all = table[0].get("kutipan", {}).get("all", 0) or 0
#             total_since = table[0].get("kutipan", {}).get("sejak_2020", 0) or 0
#         if len(table) >= 2:
#             h_all = table[1].get("indeks_h", {}).get("all", 0) or 0
#             h_since = table[1].get("indeks_h", {}).get("sejak_2020", 0) or 0
#         if len(table) >= 3:
#             i10_all = table[2].get("indeks_i10", {}).get("all", 0) or 0
#             i10_since = table[2].get("indeks_i10", {}).get("sejak_2020", 0) or 0
#     except Exception:
#         # jika struktur unexpected -> tetap aman dengan 0
#         pass

#     return total_all, total_since, h_all, h_since, i10_all, i10_since

# def extract_citation_graph(author_info):
#     """Mengembalikan DataFrame berisi kolom: 'Year' dan 'Citations'"""
#     cited_by = author_info.get("cited_by", {}) or {}
#     graph = cited_by.get("graph", []) if isinstance(cited_by.get("graph", []), list) else []
#     rows = []
#     for entry in graph:
#         # beberapa entry bisa memiliki key 'year' dan 'citations'
#         y = entry.get("year")
#         c = entry.get("citations") if "citations" in entry else entry.get("citations_count", 0)
#         if y is not None:
#             rows.append({"Year": int(y), "Citations": int(c or 0)})
#     if not rows:
#         # kosong -> kembalikan DataFrame kosong
#         return pd.DataFrame(columns=["Year", "Citations"])
#     df = pd.DataFrame(rows).sort_values("Year").reset_index(drop=True)
#     return df

# # === KONFIGURASI DAFTAR DOSEN ===
# dosen_list = [
#     {"id": "rX8zRcwAAAAJ", "nama": "Abd Hadi"},
# ]

# if __name__ == "__main__":
#     for dosen in dosen_list:
#         print(f"\n📄 Mengambil data untuk {dosen['nama']} (ID: {dosen['id']})...")
#         try:
#             author_json = get_author_info(dosen["id"])
#         except Exception as e:
#             print(f"❌ Gagal mengambil author_info untuk {dosen['nama']}: {e}")
#             continue

#         # perhatikan SerpAPI menempatkan info author di key 'author' (lihat contoh JSON)
#         author_obj = author_json.get("author") or {}

#         total_all, total_since, h_all, h_since, i10_all, i10_since = extract_citation_stats(author_json)
#         df_graph = extract_citation_graph(author_json)
#         publications = get_publications_for_author(dosen["id"])

#         # Data profil dosen (aman terhadap missing keys)
#         interests = author_obj.get("interests", []) if isinstance(author_obj.get("interests", []), list) else []
#         bidang = ", ".join([it.get("title", "") for it in interests if isinstance(it, dict)]) if interests else ""

#         df_info = pd.DataFrame([{
#             "Nama Dosen": author_obj.get("name", dosen["nama"]),
#             "Afiliasi": author_obj.get("affiliations", ""),
#             "Email": author_obj.get("email", ""),
#             "Bidang": bidang,
#             "Total Kutipan (Semua)": total_all,
#             "Total Kutipan (Sejak 2020)": total_since,
#             "h-index (Semua)": h_all,
#             "h-index (Sejak 2020)": h_since,
#             "i10-index (Semua)": i10_all,
#             "i10-index (Sejak 2020)": i10_since
#         }])

#         # Data publikasi
#         df_pub = pd.DataFrame(publications) if publications else pd.DataFrame(columns=[
#             "Judul Penelitian", "Link Publikasi", "Tipe Link", "Pengarang",
#             "Jumlah Sitasi", "Tahun Terbit", "Nama Jurnal"
#         ])

#         # Simpan ke excel dengan 3 sheet: Profil Dosen, Daftar Publikasi, Grafik Sitasi
#         file_name = f"{dosen['nama']}_publikasi.xlsx"
#         try:
#             with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
#                 df_info.to_excel(writer, sheet_name="Profil Dosen", index=False)
#                 df_pub.to_excel(writer, sheet_name="Daftar Publikasi", index=False)
#                 # sheet grafik (tahun / kutipan)
#                 df_graph.to_excel(writer, sheet_name="Grafik Sitasi", index=False)
#             print(f"✅ Data {dosen['nama']} berhasil disimpan ke file: {file_name}")
#             print(f"📊 Total Kutipan: {total_all} (Sejak2020: {total_since}) | h: {h_all}/{h_since} | i10: {i10_all}/{i10_since}")
#             if not df_graph.empty:
#                 print("📈 Grafik sitasi tersedia (sheet 'Grafik Sitasi').")
#             else:
#                 print("⚠️ Grafik sitasi kosong (tidak tersedia di JSON).")
#         except PermissionError:
#             print(f"❌ Gagal menulis file {file_name}. Pastikan file tidak sedang dibuka.")
#         except Exception as e:
#             print(f"❌ Error saat menyimpan file {file_name}: {e}")
#############################################################################################################

# YANG KE TIGA 3

import requests
import pandas as pd
import time
import openpyxl
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

# === API KEY SERPAPI ===
SERPAPI_API_KEY = "1eb991a416303d0888fee987b55b42ca60ea7fd6d75f1ee0132782515a461d36"

# === Fungsi Ambil Profil Author ===
def get_author_info(user_id):
    params = {
        "engine": "google_scholar_author",
        "author_id": user_id,
        "hl": "id",
        "api_key": SERPAPI_API_KEY
    }
    res = requests.get("https://serpapi.com/search", params=params)
    res.raise_for_status()
    return res.json()

# === Fungsi Ambil Publikasi ===
def get_publications_for_author(user_id):
    publications = []
    start = 0
    while True:
        params = {
            "engine": "google_scholar_author",
            "author_id": user_id,
            "hl": "id",
            "api_key": SERPAPI_API_KEY,
            "start": start
        }
        res = requests.get("https://serpapi.com/search", params=params)
        res.raise_for_status()
        data = res.json()
        pubs = data.get("articles", [])

        for pub in pubs:
            link = pub.get("external_link") or pub.get("link")
            publications.append({
                "Judul Penelitian": pub.get("title"),
                "Link Publikasi": link,
                "Tipe Link": "PDF" if link and isinstance(link, str) and link.lower().endswith(".pdf") else "Halaman Web",
                "Pengarang": pub.get("authors"),
                "Jumlah Sitasi": pub.get("cited_by", {}).get("value", 0) if pub.get("cited_by") else 0,
                "Tahun Terbit": pub.get("year"),
                "Nama Jurnal": pub.get("publication")
            })

        if not pubs or len(pubs) < 20:
            break
        start += len(pubs)
        time.sleep(1)
    return publications

# === Fungsi Ambil Statistik Kutipan (Semua & Sejak 2020) ===
def extract_citation_stats(author_info):
    cited_by = author_info.get("cited_by", {}) or {}
    table = cited_by.get("table", []) if isinstance(cited_by.get("table", []), list) else []

    total_all = total_since = 0
    h_all = h_since = 0
    i10_all = i10_since = 0

    try:
        if len(table) >= 1:
            total_all = table[0].get("kutipan", {}).get("all", 0) or 0
            total_since = table[0].get("kutipan", {}).get("sejak_2020", 0) or 0
        if len(table) >= 2:
            h_all = table[1].get("indeks_h", {}).get("all", 0) or 0
            h_since = table[1].get("indeks_h", {}).get("sejak_2020", 0) or 0
        if len(table) >= 3:
            i10_all = table[2].get("indeks_i10", {}).get("all", 0) or 0
            i10_since = table[2].get("indeks_i10", {}).get("sejak_2020", 0) or 0
    except Exception:
        pass

    return total_all, total_since, h_all, h_since, i10_all, i10_since

# === Fungsi Ambil Data Grafik Sitasi per Tahun ===
def extract_citation_graph(author_info):
    cited_by = author_info.get("cited_by", {}) or {}
    graph = cited_by.get("graph", []) if isinstance(cited_by.get("graph", []), list) else []
    rows = []
    for entry in graph:
        y = entry.get("year")
        c = entry.get("citations") if "citations" in entry else entry.get("citations_count", 0)
        if y is not None:
            rows.append({"Year": int(y), "Citations": int(c or 0)})
    return pd.DataFrame(rows).sort_values("Year").reset_index(drop=True) if rows else pd.DataFrame(columns=["Year", "Citations"])

# === DAFTAR DOSEN YANG AKAN DIAMBIL ===
dosen_list = [
    # {"id": "pFQaOPQAAAAJ", "nama": "Lia Farokhah"},
                            #DOSEN FTD
    # {"id": "rX8zRcwAAAAJ", "nama": "Abd Hadi"},                 #1
    # {"id": "Gdik0nIAAAAJ", "nama": "Abdul loh Eizzi Irsyada"},  #2
    # {"id": "CoAA0oAAAAAJ", "nama": "Achmad Noercholis"},        #3
    # {"id": "pCcxzLoAAAAJ", "nama": "Azwar Riza Habibi"},        #4
    # {"id": "Y3TMwyUAAAAJ", "nama": "Danang Arbian Sulistyo"},   #5
    # {"id": "pFQaOPQAAAAJ", "nama": "Lia Farokhah"},             #6
    # {"id": "QHoOc30AAAAJ", "nama": "Dr.Puji Subekti"},          #7
    # {"id": "BKfiMAgAAAAJ", "nama": "Tri Wahyuni"},              #8
    # {"id": "4O6_G7sAAAAJ", "nama": "Fadhli Almu'iini Ahda"},    #9
    # {"id": "5CjAe3QAAAAJ", "nama": "Faldi Hendrawan"},          #10
    # {"id": "D0KYQ0oAAAAJ", "nama": "Fransiska Sisilia Mukti"},  #11
    # {"id": "3ZEzr0EAAAAJ", "nama": "Handry Rochmad Dwi Happy"},   #12
    # {"id": "fpyqPioAAAAJ", "nama": "Ida Wahyuni"},              #13
    # {"id": "JxY0zO0AAAAJ", "nama": "Jaenal Arifin"},            #14
    # {"id": "iBHpmqMAAAAJ", "nama": "Lely Surya Wardani"},       #15
    {"id": "SxacVdAAAAAJ", "nama": "Lilis Widayanti"},          #16
    # {"id": "jvx-Qn0AAAAJ", "nama": "Lukman Hakim"},             #17
    # {"id": "bI3v9ZMAAAAJ", "nama": "MOHAMMAD ZAINUDDIN"},       #18
    # {"id": "5aXswucAAAAJ", "nama": "mufidatul Islamiyah"},      #19
    # {"id": "4bw9WgUAAAAJ", "nama": "Muhammad Rofiq"},           #20
    # {"id": "5jjDKXcAAAAJ", "nama": "Wayong Kabalen"},           #21
    # {"id": "O35F8U4AAAAJ", "nama": "Nur Lailatul Aqromi"},      #22
    # {"id": "ObDUmdAAAAAJ", "nama": "Philip Faster Eka "},       #23
    # {"id": "AKZfru0AAAAJ", "nama": "Rina Dewi Indahsari"},      #24
    # {"id": "XlsTPe4AAAAJ", "nama": "Samsul Arifin"},            #25
    # {"id": "A3fFhoIAAAAJ", "nama": "Sean Elbert Jeremiah"},     #26
    # {"id": "Qa0jRdsAAAAJ", "nama": "Setyorini"},                #27
    # {"id": "Q3SexakAAAAJ", "nama": "Siti Nurul Afiyah"},        #28
    # {"id": "3s3ntdwAAAAJ", "nama": "Suastika Yulia Riska"},     #29
    # {"id": "dZePwOQAAAAJ", "nama": "Sunu Jatmika"},             #30
    # {"id": "MnF3ZR4AAAAJ", "nama": "Titania Dwiandini"},        #31                              
    # {"id": "c7sKLWQAAAAJ", "nama": "Vivi Aida Fitria"},         #32
    # {"id": "_7uXJTQAAAAJ", "nama": "Yogi Widya Saka Warsaa"},   #33
    # {"id": "tt5iBlUAAAAJ", "nama": "Yudistira Sapoetra"},       #34
                            #DOSEN FEB
    # {"id": "eQPeIZEAAAAJ", "nama": "Mariana Puspa Dewi"},       #35
    # {"id": "OQYeShEAAAAJ", "nama": "MEGA MIRASAPUTRI CAHYANTI"},#36
    # {"id": "hwzcl4gAAAAJ", "nama": "MULYANINGTYAS"},            #37
    # {"id": "0IwUvkUAAAAJ", "nama": "PIPIT ROSITA ANDARSARI"},   #38
    # {"id": "IabXZ4wAAAAJ", "nama": "RISA SANTOSO"},             #39
    # {"id": "yCUmstgAAAAJ", "nama": "Widiya Dewi Anjaningrum"},  #40
    # {"id": "WgoEnSwAAAAJ", "nama": "WIDYA ADHARIYANTY RAHAYU"}, #41
    # {"id": "RAan8qcAAAAJ", "nama": "Zainul Muchlas"},           #42
    # {"id": "HCFgmC8AAAAJ", "nama": "Mohammad Bukhori"},         #43
    # {"id": "1Sq8y4kAAAAJ", "nama": "Murtianingsih Aning"},      #44
    # {"id": "Kj8hjp8AAAAJ", "nama": "Rifki Hanif"},              #45
    # {"id": "ce0Y2tMAAAAJ", "nama": "Theresia Pradiani"},        #46
    # {"id": "w3J3wEkAAAAJ", "nama": "Tin Agustina Karnawati"},   #47
    # {"id": "iBN2CwUAAAAJ", "nama": "Widi Dewi"},                #48
    # {"id": "138bEEIAAAAJ", "nama": "Yunus Handoko"},            #49
    # {"id": "bVSBxIQAAAAJ", "nama": "fadilla cahyaningtyas"},    #50
    # {"id": "MVC7XSQAAAAJ", "nama": "Hironimus Hari Kurniawan"}, #51
    # {"id": "-XwBxUgAAAAJ", "nama": "LUSSIA MARIESTI ANDRIANY"}, #53
    # {"id": "lXxg0eMAAAAJ","nama": "Abdul Aziz Muslim"},         #54
    # {"id": "300WG1AAAAAJ", "nama": "Aditya Hermawan"},          #55
    # {"id": "c7sKLWQAAAAJ", "nama": "adriani kala'lembang"},     #56
    # {"id": "iprgtdgAAAAJ", "nama": "Agus Purnomo Sidi"},        #57
    # {"id": "k6pC2rYAAAAJ", "nama": "AGUS RAHMAN ALAMSYAH"},     #58
    # {"id": "ATYgw1cAAAAJ", "nama": "AHMAD NIZAR YOGATAMA"},     #59
    # {"id": "Z0dCd9MAAAAJ", "nama": "ANIEK MURNIATI"},           #60
    # {"id": "s9uJJ_IAAAAJ", "nama": "DITYA WARDANA"},            #61
    # {"id": "kw-_iGgAAAAJ",    "nama": "Dr.FATHORRAHMAN"},         #62
    # {"id": "Uw6itsEAAAAJ", "nama": "Ike Kusdyah Rachmawati"},   #63
    # {"id": "u3nUmSEAAAAJ", "nama": "TEGUH WIDODO"},             #64
    # {"id": "ZzuMv60AAAAJ", "nama": "justita dura"},             #65
    # Tambahkan dosen lain di sini kalau mau
]

# === MAIN PROGRAM ===
if __name__ == "__main__":
    for dosen in dosen_list:
        print(f"\n📄 Mengambil data untuk {dosen['nama']} (ID: {dosen['id']})...")

        try:
            author_json = get_author_info(dosen["id"])
        except Exception as e:
            print(f"❌ Gagal mengambil author_info untuk {dosen['nama']}: {e}")
            continue

        author_obj = author_json.get("author") or {}
        total_all, total_since, h_all, h_since, i10_all, i10_since = extract_citation_stats(author_json)
        df_graph = extract_citation_graph(author_json)
        publications = get_publications_for_author(dosen["id"])

        interests = author_obj.get("interests", [])
        bidang = ", ".join([it.get("title", "") for it in interests]) if interests else ""

        df_info = pd.DataFrame([{
            "Nama Dosen": author_obj.get("name", dosen["nama"]),
            "Afiliasi": author_obj.get("affiliations", ""),
            "Email": author_obj.get("email", ""),
            "Bidang": bidang,
            "Total Kutipan (Semua)": total_all,
            "Total Kutipan (Sejak 2020)": total_since,
            "h-index (Semua)": h_all,
            "h-index (Sejak 2020)": h_since,
            "i10-index (Semua)": i10_all,
            "i10-index (Sejak 2020)": i10_since
        }])

        df_pub = pd.DataFrame(publications) if publications else pd.DataFrame(columns=[
            "Judul Penelitian", "Link Publikasi", "Tipe Link", "Pengarang",
            "Jumlah Sitasi", "Tahun Terbit", "Nama Jurnal"
        ])

        file_name = f"{dosen['nama']}_publikasi.xlsx"
        try:
            with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
                df_info.to_excel(writer, sheet_name="Profil Dosen", index=False)
                df_pub.to_excel(writer, sheet_name="Daftar Publikasi", index=False)
                df_graph.to_excel(writer, sheet_name="Grafik Sitasi", index=False)

            # === Tambahkan grafik batang ke sheet "Grafik Sitasi" ===
            wb = load_workbook(file_name)
            ws = wb["Grafik Sitasi"]

            if not df_graph.empty:
                chart = BarChart()
                chart.title = "Grafik Sitasi per Tahun"
                chart.x_axis.title = "Tahun"
                chart.y_axis.title = "Jumlah Sitasi"

                data = Reference(ws, min_col=2, min_row=1, max_row=len(df_graph) + 1)
                cats = Reference(ws, min_col=1, min_row=2, max_row=len(df_graph) + 1)
                chart.add_data(data, titles_from_data=True)
                chart.set_categories(cats)

                ws.add_chart(chart, "D5")
                wb.save(file_name)

            print(f"✅ Data {dosen['nama']} berhasil disimpan ke file: {file_name}")
            print(f"📊 Total Kutipan: {total_all} (Sejak2020: {total_since}) | h: {h_all}/{h_since} | i10: {i10_all}/{i10_since}")
            if not df_graph.empty:
                print("📈 Grafik sitasi berhasil ditambahkan ke Excel.")
            else:
                print("⚠️ Tidak ada data grafik sitasi.")

        except PermissionError:
            print(f"❌ File {file_name} sedang dibuka, tutup dulu lalu jalankan ulang.")
        except Exception as e:
            print(f"❌ Error saat menyimpan file {file_name}: {e}")
